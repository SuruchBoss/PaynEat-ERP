"""Generate the draft ERP import templates used to test ExcelToGo (ADR-0016).

Draft 0 — a design sample, not a contract. The real template format is fixed by the ERP ticket
that builds the importer. All data is the fictional fried-chicken chain; nothing refers to a real
company.

    python -m venv .venv && .venv/bin/pip install openpyxl==3.1.5
    .venv/bin/python make_sample_templates.py
"""

from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

UNITS = ["kg", "g", "pcs", "L", "ml", "case", "bag", "tin"]
LOCATIONS = [
    ("PLANT-01", "โรงงานบางนา", "Bang Na plant", "plant"),
    ("BR-SILOM", "สาขาสีลม", "Silom branch", "branch"),
    ("BR-ARI", "สาขาอารีย์", "Ari branch", "branch"),
    ("BR-BANGNA", "สาขาบางนา", "Bang Na branch", "branch"),
]
# code, name_th, name_en, base unit, variable weight, shelf life days, purchase unit, factor
ITEMS = [
    ("WHOLE-CHICKEN", "ไก่ทั้งตัว", "Whole chicken", "kg", "TRUE", 5, "case", 20),
    ("CHICKEN-BREAST", "อกไก่", "Chicken breast", "pcs", "FALSE", 4, "pcs", 1),
    ("CHICKEN-THIGH", "สะโพกไก่", "Chicken thigh", "pcs", "FALSE", 4, "pcs", 1),
    ("CHICKEN-DRUMSTICK", "น่องไก่", "Chicken drumstick", "pcs", "FALSE", 4, "pcs", 1),
    ("CHICKEN-WING", "ปีกไก่", "Chicken wing", "pcs", "FALSE", 4, "pcs", 1),
    ("CHICKEN-FRAME", "โครงไก่", "Chicken frame", "kg", "TRUE", 3, "kg", 1),
    ("FLOUR", "แป้งชุบทอด", "Batter flour", "kg", "FALSE", 180, "bag", 25),
    ("FRYING-OIL", "น้ำมันทอด", "Frying oil", "L", "FALSE", 365, "tin", 18),
    ("SEASONING", "ผงปรุงรส", "Seasoning", "kg", "FALSE", 365, "bag", 5),
]

HEADER = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5597")


def header(ws, names):
    ws.append(names)
    for cell in ws[1]:
        cell.font = HEADER
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A2"
    for i, name in enumerate(names, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(name) + 4)


def unlock(ws, first_row, last_row, last_col):
    for row in ws.iter_rows(min_row=first_row, max_row=last_row, max_col=last_col):
        for cell in row:
            cell.protection = Protection(locked=False)


def list_rule(ws, formula, sqref, prompt):
    rule = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True)
    rule.error = "Choose a value from the list."
    rule.prompt = prompt
    rule.add(sqref)
    ws.add_data_validation(rule)


def build(path, item_codes, balance_rows):
    wb = Workbook()

    readme = wb.active
    readme.title = "README"
    for line in [
        "PaynEat ERP import template — DRAFT 0, a design sample, not a contract (ADR-0016).",
        "แม่แบบนำเข้าข้อมูล PaynEat ERP — ฉบับร่าง 0 ใช้ทดสอบการออกแบบ ยังไม่ใช่สัญญา",
        "Fill the unlocked cells in Items, Locations and OpeningBalance. Lists come from the hidden, locked Ref sheet,",
        "which the ERP generates from its current master data. The ERP re-validates every row on import.",
        "All data is the fictional fried-chicken chain.",
    ]:
        readme.append([line])
    readme.column_dimensions["A"].width = 110

    ref = wb.create_sheet("Ref")
    ref.append(["unit", "location_code", "item_code"])
    for i in range(max(len(UNITS), len(LOCATIONS), len(item_codes))):
        ref.append([
            UNITS[i] if i < len(UNITS) else None,
            LOCATIONS[i][0] if i < len(LOCATIONS) else None,
            item_codes[i] if i < len(item_codes) else None,
        ])
    units_ref = f"Ref!$A$2:$A${len(UNITS) + 1}"
    locations_ref = f"Ref!$B$2:$B${len(LOCATIONS) + 1}"
    items_ref = f"Ref!$C$2:$C${len(item_codes) + 1}"
    ref.sheet_state = "hidden"
    ref.protection.sheet = True

    items = wb.create_sheet("Items")
    header(items, ["item_code", "name_th", "name_en", "base_unit", "variable_weight",
                   "shelf_life_days", "purchase_unit", "purchase_factor"])
    for row in ITEMS:
        items.append(list(row))
    last = 200
    unlock(items, 2, last, 8)
    list_rule(items, units_ref, f"D2:D{last}", "Base unit")
    list_rule(items, '"TRUE,FALSE"', f"E2:E{last}", "Weight and piece count both recorded?")
    list_rule(items, units_ref, f"G2:G{last}", "Unit the supplier sells in")
    whole = DataValidation(type="whole", operator="between", formula1="0", formula2="3650", allow_blank=True)
    whole.add(f"F2:F{last}")
    items.add_data_validation(whole)
    factor = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True)
    factor.add(f"H2:H{last}")
    items.add_data_validation(factor)
    code = DataValidation(type="textLength", operator="between", formula1="2", formula2="32", allow_blank=True)
    code.add(f"A2:A{last}")
    items.add_data_validation(code)
    items.protection.sheet = True

    locations = wb.create_sheet("Locations")
    header(locations, ["location_code", "name_th", "name_en", "type"])
    for row in LOCATIONS:
        locations.append(list(row))
    unlock(locations, 2, 50, 4)
    list_rule(locations, '"plant,warehouse,branch"', "D2:D50", "Location type")
    loc_code = DataValidation(type="textLength", operator="between", formula1="2", formula2="32", allow_blank=True)
    loc_code.add("A2:A50")
    locations.add_data_validation(loc_code)
    locations.protection.sheet = True

    balances = wb.create_sheet("OpeningBalance")
    header(balances, ["location_code", "item_code", "quantity", "secondary_quantity", "unit_cost", "expiry_date"])
    for row in balance_rows:
        balances.append(row)
    last_b = max(len(balance_rows) + 1, 200)
    unlock(balances, 2, last_b, 6)
    list_rule(balances, locations_ref, f"A2:A{last_b}", "Location code, exactly as in the ERP")
    list_rule(balances, items_ref, f"B2:B{last_b}", "Item code, exactly as in the ERP")
    for col, op in (("C", "greaterThan"), ("D", "greaterThanOrEqual"), ("E", "greaterThanOrEqual")):
        rule = DataValidation(type="decimal", operator=op, formula1="0", allow_blank=True)
        rule.add(f"{col}2:{col}{last_b}")
        balances.add_data_validation(rule)
    expiry = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True)
    expiry.add(f"F2:F{last_b}")
    balances.add_data_validation(expiry)
    for row in balances.iter_rows(min_row=2, min_col=6, max_col=6):
        row[0].number_format = "yyyy-mm-dd"
    balances.protection.sheet = True

    wb.active = wb.sheetnames.index("Items")
    wb.save(path)


def main():
    real_codes = [row[0] for row in ITEMS]
    today = date(2026, 10, 1)
    small_rows = [
        ["PLANT-01", "WHOLE-CHICKEN", 432.0, 240, 78.5, today + timedelta(days=4)],
        ["PLANT-01", "FLOUR", 250.0, None, 32.0, today + timedelta(days=150)],
        ["PLANT-01", "FRYING-OIL", 180.0, None, 54.0, today + timedelta(days=300)],
        ["BR-SILOM", "CHICKEN-DRUMSTICK", 120.0, None, 11.2, today + timedelta(days=3)],
    ]
    build("sample-import-template.xlsx", real_codes, small_rows)

    # Size test: 3,000 synthetic item codes in Ref and 3,000 opening-balance rows with dropdowns.
    synthetic = real_codes + [f"SYN-{n:05d}" for n in range(1, 3000 - len(real_codes) + 1)]
    large_rows = [
        [LOCATIONS[n % len(LOCATIONS)][0], synthetic[n], 10.0 + n % 7, None, 20.0 + n % 13,
         today + timedelta(days=30 + n % 90)]
        for n in range(3000)
    ]
    build("sample-import-template-large.xlsx", synthetic, large_rows)


if __name__ == "__main__":
    main()
